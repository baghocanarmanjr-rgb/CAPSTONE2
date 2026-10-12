"""Offline regression test for PPMP extraction across supported document types.
Does not touch the database. It verifies extraction + mapping before human review.
"""
from pathlib import Path
from tempfile import TemporaryDirectory
import csv
from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from docx import Document
import fitz

from aoq_flexible import extract_document
from offline_ocr_mapper import map_document, PPMP_FIELDS

FIELDS = PPMP_FIELDS
TEXT = """Project Particulars:\nAcquisition of Network Infrastructure Equipment\nClassification:\nGoods\nRequired Units / Technical Specs:\n4 units Managed Gigabit Switch, 24-port, rack-mountable, VLAN support\nProcurement Method:\nSmall Value Procurement\nPre-Bid Requirement:\nNot Applicable\nProcurement Start:\nApril 2026\nCompletion Target:\nMay 2026\nDelivery Schedule:\nWithin 30 calendar days after notice to proceed\nFunding Source:\nGeneral Fund\nApproved Budget for the Contract:\nPHP 80,000.00\nAttachments:\nTechnical Specifications; Market Survey\nAdditional Notes:\nFor municipal network upgrade and redundancy\n"""

EXPECTED = {
    FIELDS['description']['label']: 'Acquisition of Network Infrastructure Equipment',
    FIELDS['project_type']['label']: 'Goods',
    FIELDS['quantity_size']['label']: '4 units Managed Gigabit Switch, 24-port, rack-mountable, VLAN support',
    FIELDS['procurement_mode']['label']: 'Small Value Procurement',
    FIELDS['pre_procurement']['label']: 'Not Applicable',
    FIELDS['start_date']['label']: 'April 2026',
    FIELDS['end_date']['label']: 'May 2026',
    FIELDS['delivery_period']['label']: 'Within 30 calendar days after notice to proceed',
    FIELDS['fund_source']['label']: 'General Fund',
    FIELDS['budget']['label']: 'PHP 80,000.00',
    FIELDS['supporting_documents']['label']: 'Technical Specifications; Market Survey',
    FIELDS['remarks']['label']: 'For municipal network upgrade and redundancy',
}

KV_ROWS = [
    ['Project Particulars', EXPECTED[FIELDS['description']['label']]],
    ['Classification', 'Goods'],
    ['Required Units / Technical Specs', EXPECTED[FIELDS['quantity_size']['label']]],
    ['Procurement Method', 'Small Value Procurement'],
    ['Pre-Bid Requirement', 'Not Applicable'],
    ['Procurement Start', 'April 2026'],
    ['Completion Target', 'May 2026'],
    ['Delivery Schedule', EXPECTED[FIELDS['delivery_period']['label']]],
    ['Funding Source', 'General Fund'],
    ['Approved Budget for the Contract', 'PHP 80,000.00'],
    ['Attachments', EXPECTED[FIELDS['supporting_documents']['label']]],
    ['Additional Notes', EXPECTED[FIELDS['remarks']['label']]],
]


def assert_mapping(path: Path, label: str, min_exact: int = 10):
    doc = extract_document(str(path))
    mapped = map_document(doc)
    assert mapped['rows'], f'{label}: no mapped rows'
    fields = mapped['rows'][0]['fields']
    exact = sum(1 for k, v in EXPECTED.items() if fields.get(k) == v)
    # OCR images can have minor punctuation differences. Native formats should be exact.
    assert exact >= min_exact, f'{label}: only {exact}/12 exact fields: {fields}'
    assert fields.get(FIELDS['description']['label']), f'{label}: description missing'
    assert fields.get(FIELDS['budget']['label']), f'{label}: budget missing'
    print(f'PASS {label}: {exact}/12 exact fields; {doc["extraction_method"]}')


def make_image(path: Path, fmt: str):
    img = Image.new('RGB', (2400, 2600), 'white')
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype('DejaVuSans.ttf', 38)
    except Exception:
        font = ImageFont.load_default()
    y = 70
    for line in TEXT.splitlines():
        draw.text((90, y), line, fill='black', font=font)
        y += 62
    img.save(path, format=fmt, quality=96 if fmt == 'JPEG' else None, dpi=(300, 300))


def main():
    with TemporaryDirectory() as tmp:
        d = Path(tmp)
        # TXT
        (d/'sample.txt').write_text(TEXT, encoding='utf-8')
        assert_mapping(d/'sample.txt', 'TXT', 12)

        # CSV key/value form
        with (d/'sample.csv').open('w', encoding='utf-8-sig', newline='') as fh:
            csv.writer(fh).writerows(KV_ROWS)
        assert_mapping(d/'sample.csv', 'CSV', 12)

        # XLSX key/value form
        wb = Workbook(); ws = wb.active; ws.title = 'PPMP'
        for row in KV_ROWS: ws.append(row)
        wb.save(d/'sample.xlsx')
        assert_mapping(d/'sample.xlsx', 'XLSX', 12)
        wb.save(d/'sample.xlsm')
        assert_mapping(d/'sample.xlsm', 'XLSM', 12)

        # DOCX key/value form
        docx = Document(); table = docx.add_table(rows=0, cols=2)
        for left, right in KV_ROWS:
            cells = table.add_row().cells; cells[0].text = left; cells[1].text = right
        docx.save(d/'sample.docx')
        assert_mapping(d/'sample.docx', 'DOCX', 12)

        # Native PDF
        pdf = fitz.open(); page = pdf.new_page(width=612, height=1000)
        page.insert_textbox(fitz.Rect(40, 40, 570, 960), TEXT, fontsize=10)
        pdf.save(d/'native.pdf'); pdf.close()
        assert_mapping(d/'native.pdf', 'native PDF', 12)

        # Image OCR
        make_image(d/'sample.png', 'PNG')
        assert_mapping(d/'sample.png', 'PNG OCR', 9)
        make_image(d/'sample.jpg', 'JPEG')
        assert_mapping(d/'sample.jpg', 'JPG OCR', 9)
        make_image(d/'sample.bmp', 'BMP')
        assert_mapping(d/'sample.bmp', 'BMP OCR', 9)
        make_image(d/'sample.tiff', 'TIFF')
        assert_mapping(d/'sample.tiff', 'TIFF OCR', 9)

        # Scanned PDF: embed image only, forcing OCR.
        scan_pdf = fitz.open(); page = scan_pdf.new_page(width=612, height=792)
        page.insert_image(page.rect, filename=str(d/'sample.png'))
        scan_pdf.save(d/'scanned.pdf'); scan_pdf.close()
        assert_mapping(d/'scanned.pdf', 'scanned PDF OCR', 9)

    print('PASS: supported-format extraction regression complete')

if __name__ == '__main__':
    main()
